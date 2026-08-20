from __future__ import annotations

import base64
import re
from collections.abc import Iterable
from io import BytesIO
from pathlib import Path
from typing import TypedDict

import xlrd
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from budget_agent.models import BudgetSource, UploadedDocument

SUPPORTED_EXTENSIONS = {".pdf", ".doc", ".docx", ".xls", ".xlsx"}
REQUIRED_SOURCES: tuple[BudgetSource, ...] = (
    "project_summary",
    "per_budget",
    "qs_estimate",
    "vendor_quotation",
    "mapping_rules",
    "historical_rates",
)
SOURCE_LABELS: dict[BudgetSource, str] = {
    "project_summary": "Project Summary (project name and construction area)",
    "per_budget": "PER Budget",
    "qs_estimate": "QS Estimate",
    "vendor_quotation": "Vendor Quotations",
    "mapping_rules": "Mapping Rules",
    "historical_rates": "Historical Unit Rates",
}
SOURCE_FILE_PATTERNS: dict[BudgetSource, tuple[str, ...]] = {
    "project_summary": ("project summary", "project brief", "general information"),
    "per_budget": ("per budget", "perbudget", "项目预算"),
    "qs_estimate": ("qs estimation", "qs estimate", "quantity surveyor estimate", "qs 估算"),
    "vendor_quotation": ("vendor quotation", "vendor quote", "tender quotation", "供应商报价"),
    "mapping_rules": ("mapping rules", "mapping rule", "wbs mapping", "映射规则"),
    "historical_rates": (
        "historical unit rates",
        "historical rates",
        "historic unit rate",
        "history",
        "历史单价",
    ),
}


class DocumentValidationError(ValueError):
    pass


class DocumentChunk(TypedDict):
    filename: str
    chunk_id: str
    content: str


def decode_document(document: UploadedDocument) -> tuple[str, bytes]:
    filename = Path(document.filename).name
    extension = Path(filename).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentValidationError(f"{filename}: unsupported file type. Supported types: {supported}.")
    try:
        content = base64.b64decode(document.file_base64, validate=True)
    except ValueError as exc:
        raise DocumentValidationError(f"{filename}: file content is not valid Base64.") from exc
    if not content:
        raise DocumentValidationError(f"{filename}: file is empty.")
    return filename, content


def _xlsx_text(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    sections: list[str] = []
    for worksheet in workbook.worksheets:
        sections.append(f"### SHEET: {worksheet.title}")
        for row in worksheet.iter_rows(values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row]
            if any(values):
                sections.append("\t".join(values))
    return "\n".join(sections)


def _xls_text(content: bytes) -> str:
    workbook = xlrd.open_workbook(file_contents=content)
    sections: list[str] = []
    for worksheet in workbook.sheets():
        sections.append(f"### SHEET: {worksheet.name}")
        for row_index in range(worksheet.nrows):
            values = [str(value).strip() for value in worksheet.row_values(row_index)]
            if any(values):
                sections.append("\t".join(values))
    return "\n".join(sections)


def _docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    sections = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        sections.append("### TABLE")
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(values):
                sections.append("\t".join(values))
    return "\n".join(sections)


def _pdf_text(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    return "\n".join(
        f"### PAGE {index}\n{page.extract_text() or ''}"
        for index, page in enumerate(reader.pages, start=1)
    )


def _legacy_doc_text(content: bytes) -> str:
    ascii_runs = re.findall(rb"[\x20-\x7e]{4,}", content)
    utf16_runs = re.findall(rb"(?:[\x20-\x7e]\x00){4,}", content)
    parts = [run.decode("cp1252", errors="ignore") for run in ascii_runs]
    parts.extend(run.decode("utf-16le", errors="ignore") for run in utf16_runs)
    return "\n".join(dict.fromkeys(part.strip() for part in parts if part.strip()))


def extract_document_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.casefold()
    try:
        if extension == ".xlsx":
            text = _xlsx_text(content)
        elif extension == ".xls":
            text = _xls_text(content)
        elif extension == ".docx":
            text = _docx_text(content)
        elif extension == ".doc":
            text = _legacy_doc_text(content)
        elif extension == ".pdf":
            text = _pdf_text(content)
        else:
            raise DocumentValidationError(f"{filename}: unsupported file type.")
    except DocumentValidationError:
        raise
    except Exception as exc:
        raise DocumentValidationError(f"{filename}: unable to read document: {exc}") from exc
    if not text.strip():
        raise DocumentValidationError(f"{filename}: no readable text or table data was found.")
    return text


def _normalized(value: str) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", " ", value.casefold()).strip()


def _has_any(value: str, patterns: tuple[str, ...]) -> bool:
    return any(_normalized(pattern) in value for pattern in patterns)


def identify_sources(filename: str, text: str) -> set[BudgetSource]:
    stem = _normalized(Path(filename).stem)
    content = _normalized(text[:120000])
    sources = {
        source
        for source, patterns in SOURCE_FILE_PATTERNS.items()
        if _has_any(stem, patterns)
    }

    if _has_any(
        content,
        ("construction area", "project summary", "project brief", "项目概况", "施工面积"),
    ):
        sources.add("project_summary")
    if (
        "sheet per budget" in content
        or "per category budget amount" in content
        or ("per category" in content and "budget amount" in content)
    ):
        sources.add("per_budget")
    if (
        "sheet qs estimation" in content
        or "qs estimate category total amount" in content
        or ("qs estimate category" in content and "total amount" in content)
    ):
        sources.add("qs_estimate")
    if (
        "sheet vendor quotation" in content
        or ("description" in content and "gc total" in content and "vendor" in content)
    ):
        sources.add("vendor_quotation")
    if (
        "sheet mapping rules" in content
        or "vendor wbs description" in content
        or "wbs mapping" in content
    ):
        sources.add("mapping_rules")
    if (
        "sheet historical unit rates" in content
        or "historical unit rate" in content
        or ("qs category" in content and "usd sqm" in content and "history" in stem)
    ):
        sources.add("historical_rates")
    return sources


def prepare_documents(
    documents: Iterable[UploadedDocument],
) -> tuple[list[dict[str, str]], set[BudgetSource], bytes]:
    extracted: list[dict[str, str]] = []
    recognized: set[BudgetSource] = set()
    digest_parts: list[bytes] = []
    for document in documents:
        filename, content = decode_document(document)
        text = extract_document_text(filename, content)
        sources = identify_sources(filename, text)
        recognized.update(sources)
        extracted.append(
            {
                "filename": filename,
                "recognized_sources": ", ".join(sorted(sources)) or "unclassified",
                "content": text[:120000],
            }
        )
        digest_parts.extend((filename.encode(), b"\0", content, b"\0"))
    if not extracted:
        raise DocumentValidationError("Upload at least one budget source file.")
    return extracted, recognized, b"".join(digest_parts)


def chunk_documents(
    documents: Iterable[dict[str, str]],
    max_chars: int = 3000,
    overlap_lines: int = 3,
) -> list[DocumentChunk]:
    chunks: list[DocumentChunk] = []
    for document in documents:
        filename = document["filename"]
        lines = [line.strip() for line in document["content"].splitlines() if line.strip()]
        start = 0
        chunk_number = 1
        while start < len(lines):
            end = start
            size = 0
            while end < len(lines):
                next_size = len(lines[end]) + 1
                if end > start and size + next_size > max_chars:
                    break
                size += next_size
                end += 1
            chunks.append(
                {
                    "filename": filename,
                    "chunk_id": f"{filename}#{chunk_number}",
                    "content": "\n".join(lines[start:end]),
                }
            )
            if end >= len(lines):
                break
            start = max(start + 1, end - overlap_lines)
            chunk_number += 1
    return chunks


def _search_tokens(value: str) -> set[str]:
    normalized = value.casefold()
    words = set(re.findall(r"[a-z0-9]+(?:[._/%-][a-z0-9]+)*", normalized))
    chinese_runs = re.findall(r"[\u4e00-\u9fff]+", normalized)
    for run in chinese_runs:
        words.update(run[index : index + 2] for index in range(max(1, len(run) - 1)))
    return {token for token in words if len(token) > 1 or token.isdigit()}


def retrieve_document_chunks(
    chunks: Iterable[DocumentChunk],
    question: str,
    filename: str | None = None,
    limit: int = 10,
) -> list[DocumentChunk]:
    available = list(chunks)
    if filename:
        requested = filename.casefold()
        available = [chunk for chunk in available if chunk["filename"].casefold() == requested]
        if not available:
            raise DocumentValidationError(f"Document not found in this report: {filename}.")
    if not available:
        return []

    query_tokens = _search_tokens(question)
    query_numbers = {token for token in query_tokens if any(char.isdigit() for char in token)}
    scored: list[tuple[float, int, DocumentChunk]] = []
    for index, chunk in enumerate(available):
        content_tokens = _search_tokens(f"{chunk['filename']} {chunk['content']}")
        overlap = query_tokens.intersection(content_tokens)
        score = float(len(overlap))
        score += 2.0 * len(query_numbers.intersection(content_tokens))
        score += sum(min(len(token), 12) / 12 for token in overlap)
        if question.casefold() in chunk["content"].casefold():
            score += 8.0
        scored.append((score, -index, chunk))

    scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
    matched = [item[2] for item in scored if item[0] > 0][:limit]
    if matched:
        return matched
    return [item[2] for item in scored[: min(2, limit)]]


def source_status(recognized: set[BudgetSource]) -> dict:
    missing = [source for source in REQUIRED_SOURCES if source not in recognized]
    return {
        "complete": not missing,
        "recognized_sources": [
            {"id": source, "name": SOURCE_LABELS[source]}
            for source in REQUIRED_SOURCES
            if source in recognized
        ],
        "missing_sources": [
            {"id": source, "name": SOURCE_LABELS[source]}
            for source in missing
        ],
    }

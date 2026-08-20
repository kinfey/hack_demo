from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from budget_agent.models import (
    BudgetEvaluation,
    Comparison,
    Decision,
    FieldComparison,
    QSRequirementRow,
    ScopeComparison,
    StructuredBudgetInput,
    TrafficLight,
    VendorAssessment,
    VendorQuoteRow,
)

REQUIRED_SHEETS = {
    "Summary",
    "PER Budget",
    "QS Estimation",
    "Vendor Quotation",
    "MAPPING RULES",
    "Historical Unit Rates",
}


class WorkbookValidationError(ValueError):
    pass


def _text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _key(value: Any) -> str:
    return _text(value).casefold()


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int | float):
        return float(value)
    try:
        return float(str(value).replace(",", "").replace("$", "").strip())
    except ValueError:
        return None


def _similarity(left: str, right: str) -> float:
    left_key = _key(left)
    right_key = _key(right)
    if not left_key or not right_key:
        return 0.0
    if left_key == right_key:
        return 1.0
    left_tokens = set(re.findall(r"[a-z0-9\u4e00-\u9fff]+", left_key))
    right_tokens = set(re.findall(r"[a-z0-9\u4e00-\u9fff]+", right_key))
    token_score = (
        len(left_tokens & right_tokens) / len(left_tokens | right_tokens)
        if left_tokens and right_tokens
        else 0.0
    )
    return max(SequenceMatcher(None, left_key, right_key).ratio(), token_score)


def _text_field(expected: str, offered: str) -> FieldComparison:
    if not expected or not offered:
        return FieldComparison(
            expected=expected or "Not specified",
            offered=offered or "Not provided",
            status="missing",
            note="Confirm this field before commercial alignment.",
        )
    similarity = _similarity(expected, offered)
    if similarity >= 0.82:
        status, note = "match", "Vendor response is aligned with the QS requirement."
    elif similarity >= 0.48 or _key(expected) in _key(offered) or _key(offered) in _key(expected):
        status, note = "review", "Wording differs; obtain a clause-by-clause compliance confirmation."
    else:
        status, note = "mismatch", "Material difference detected; resolve before award."
    return FieldComparison(expected=expected, offered=offered, status=status, note=note)


def _quantity_field(qs_row: QSRequirementRow, vendor_row: Any) -> FieldComparison:
    expected = (
        f"{qs_row.quantity:g} {qs_row.unit}".strip()
        if qs_row.quantity is not None
        else qs_row.unit or "Not specified"
    )
    offered = (
        f"{vendor_row.quantity:g} {vendor_row.unit}".strip()
        if vendor_row.quantity is not None
        else vendor_row.unit or "Not provided"
    )
    if (
        qs_row.quantity is None
        or qs_row.quantity <= 0
        or vendor_row.quantity is None
        or not qs_row.unit
        or not vendor_row.unit
    ):
        return FieldComparison(
            expected=expected,
            offered=offered,
            status="missing",
            note="Confirm both numeric quantity and measurement unit.",
        )
    if _key(qs_row.unit) != _key(vendor_row.unit):
        return FieldComparison(
            expected=expected,
            offered=offered,
            status="mismatch",
            note="Measurement units differ; normalize units and reprice before comparison.",
        )
    variance = vendor_row.quantity / qs_row.quantity - 1 if qs_row.quantity else 0.0
    if abs(variance) <= 0.05:
        status, note = "match", f"Quantity variance is {variance:+.1%}, within ±5%."
    elif abs(variance) <= 0.15:
        status, note = "review", f"Quantity variance is {variance:+.1%}; validate take-off assumptions."
    else:
        status, note = "mismatch", f"Quantity variance is {variance:+.1%}; reconcile the take-off."
    return FieldComparison(expected=expected, offered=offered, status=status, note=note)


def _scope_comparisons(data: StructuredBudgetInput) -> list[ScopeComparison]:
    if not data.qs_rows:
        return []
    results: list[ScopeComparison] = []
    mappings = {_key(source): target for source, target in data.mappings.items()}
    status_rank = {"match": 0, "missing": 1, "review": 2, "mismatch": 3}
    color_by_rank: dict[int, TrafficLight] = {0: "green", 1: "gray", 2: "yellow", 3: "red"}
    for vendor_row in data.vendor_rows:
        if not any(
            (
                vendor_row.construction_scope,
                vendor_row.specification,
                vendor_row.brand,
                vendor_row.quantity is not None,
                vendor_row.unit,
            )
        ):
            continue
        mapped_category = mappings.get(_key(vendor_row.description), "")
        qs_row = max(
            data.qs_rows,
            key=lambda row: max(
                _similarity(vendor_row.description, row.description),
                _similarity(mapped_category, row.description),
                _similarity(vendor_row.construction_scope, row.construction_scope),
            ),
        )
        match_score = max(
            _similarity(vendor_row.description, qs_row.description),
            _similarity(mapped_category, qs_row.description),
            _similarity(vendor_row.construction_scope, qs_row.construction_scope),
        )
        vendors = [vendor_row.vendor] if vendor_row.vendor else list(vendor_row.amounts_usd) or ["Vendor"]
        if match_score < 0.30:
            fields = [
                FieldComparison(
                    expected="QS line not identified",
                    offered=value or "Not provided",
                    status="missing",
                    note="Map this vendor item to a QS requirement before award.",
                )
                for value in (
                    vendor_row.construction_scope,
                    vendor_row.specification,
                    vendor_row.brand,
                    (
                        f"{vendor_row.quantity:g} {vendor_row.unit}".strip()
                        if vendor_row.quantity is not None
                        else vendor_row.unit
                    ),
                )
            ]
            qs_item: str | None = None
        else:
            fields = [
                _text_field(qs_row.construction_scope, vendor_row.construction_scope),
                _text_field(qs_row.specification, vendor_row.specification),
                _text_field(qs_row.brand, vendor_row.brand),
                _quantity_field(qs_row, vendor_row),
            ]
            qs_item = qs_row.description
        worst = max(status_rank[field.status] for field in fields)
        issues = [
            name
            for name, field in zip(
                ("construction scope", "specification", "brand", "quantity/unit"),
                fields,
                strict=True,
            )
            if field.status != "match"
        ]
        recommendation = (
            "Accept technical alignment; retain the stated requirements in the contract."
            if not issues
            else f"Resolve {', '.join(issues)} through a written compliance schedule before award."
        )
        for vendor in vendors:
            results.append(
                ScopeComparison(
                    vendor=vendor or "Vendor",
                    vendor_item=vendor_row.description,
                    qs_item=qs_item,
                    construction_scope=fields[0],
                    specification=fields[1],
                    brand=fields[2],
                    quantity_unit=fields[3],
                    overall_status=color_by_rank[worst],
                    recommendation=recommendation,
                )
            )
    return results


def _find_label_number(ws: Worksheet, label: str) -> float | None:
    target = _key(label)
    for row in ws.iter_rows():
        for cell in row:
            if target in _key(cell.value):
                for candidate in row[cell.column :]:
                    number = _number(candidate.value)
                    if number is not None:
                        return number
    return None


def _two_column_table(ws: Worksheet, *, total_prefixes: tuple[str, ...]) -> dict[str, float]:
    result: dict[str, float] = {}
    for row in ws.iter_rows(values_only=True):
        name = _text(row[0] if row else None)
        amount = _number(row[1] if len(row) > 1 else None)
        if not name or amount is None or _key(name).startswith(total_prefixes):
            continue
        result[name] = amount
    return result


def _mapping_table(ws: Worksheet) -> dict[str, str]:
    result: dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        source = _text(row[0] if row else None)
        target = _text(row[1] if len(row) > 1 else None)
        if source and target and "vendor wbs" not in _key(source):
            result[_key(source)] = target
    return result


TECHNICAL_HEADER_ALIASES = {
    "description": ("description", "item description", "qs estimate category", "项目描述", "工作内容"),
    "construction_scope": ("construction scope", "scope of work", "work scope", "施工范围"),
    "specification": ("specification", "specifications", "spec", "规格"),
    "brand": ("brand", "make", "manufacturer", "品牌", "厂家"),
    "quantity": ("quantity", "qty", "数量"),
    "unit": ("unit", "uom", "u/m", "单位"),
}


def _technical_columns(ws: Worksheet) -> tuple[int, dict[str, int]] | None:
    best: tuple[int, dict[str, int]] | None = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        columns: dict[str, int] = {}
        for cell in row:
            label = _key(cell.value)
            for field, aliases in TECHNICAL_HEADER_ALIASES.items():
                if any(alias == label or alias in label for alias in aliases):
                    columns.setdefault(field, cell.column)
        if "description" in columns and len(columns) >= 2:
            candidate = (row[0].row, columns)
            if best is None or len(columns) > len(best[1]):
                best = candidate
    return best


def _quantity_and_unit(quantity_value: Any, unit_value: Any) -> tuple[float | None, str]:
    quantity = _number(quantity_value)
    unit = _text(unit_value)
    if quantity is None and quantity_value is not None:
        combined = _text(quantity_value)
        match = re.match(r"^\s*([\d,.]+)\s*(.*)$", combined)
        if match:
            quantity = _number(match.group(1))
            unit = unit or match.group(2).strip()
    return quantity, unit


def _qs_requirement_rows(ws: Worksheet) -> list[QSRequirementRow]:
    header = _technical_columns(ws)
    if not header:
        return []
    header_row, columns = header
    technical_fields = ("construction_scope", "specification", "brand", "quantity", "unit")
    if not any(field in columns for field in technical_fields):
        return []
    rows: list[QSRequirementRow] = []
    for row_index in range(header_row + 1, ws.max_row + 1):
        description = _text(ws.cell(row_index, columns["description"]).value)
        if not description or _key(description).startswith("total"):
            continue
        quantity_value = ws.cell(row_index, columns["quantity"]).value if "quantity" in columns else None
        unit_value = ws.cell(row_index, columns["unit"]).value if "unit" in columns else None
        quantity, unit = _quantity_and_unit(quantity_value, unit_value)
        rows.append(
            QSRequirementRow(
                description=description,
                construction_scope=(
                    _text(ws.cell(row_index, columns["construction_scope"]).value)
                    if "construction_scope" in columns
                    else ""
                ),
                specification=(
                    _text(ws.cell(row_index, columns["specification"]).value)
                    if "specification" in columns
                    else ""
                ),
                brand=_text(ws.cell(row_index, columns["brand"]).value) if "brand" in columns else "",
                quantity=quantity,
                unit=unit,
            )
        )
    return rows


def _decision(variance: float, reasonable: float, review: float) -> Decision:
    absolute = abs(variance)
    if absolute <= reasonable:
        return "reasonable"
    if absolute <= review:
        return "review_required"
    return "significant_concern"


def _total_decision(variance: float) -> Decision:
    absolute = abs(variance)
    if absolute <= 0.10:
        return "approve"
    if absolute <= 0.20:
        return "conditional_approval"
    return "reject"


def _status_color(decision: Decision) -> TrafficLight:
    if decision in {"approve", "reasonable"}:
        return "green"
    if decision in {"conditional_approval", "review_required"}:
        return "yellow"
    if decision in {"reject", "significant_concern"}:
        return "red"
    return "gray"


def _vendor_table(
    ws: Worksheet,
    mappings: dict[str, str],
) -> tuple[dict[str, dict[str, float]], dict[str, float], list[str], list[VendorQuoteRow]]:
    header_row = None
    vendor_columns: dict[int, str] = {}
    description_column = None
    for row in ws.iter_rows():
        labels = [_text(cell.value) for cell in row]
        if any(label.casefold().startswith("vendor-") for label in labels):
            header_row = row[0].row
            vendor_columns = {
                cell.column: _text(cell.value).replace("Vendor-", "Vendor ")
                for cell in row
                if _key(cell.value).startswith("vendor-")
            }
        if any("description" in label.casefold() for label in labels):
            description_column = next(
                cell.column for cell in row if "description" in _key(cell.value)
            )
        if header_row and description_column:
            break
    if not header_row or not vendor_columns or not description_column:
        raise WorkbookValidationError("Vendor Quotation headers could not be identified.")

    category_totals: dict[str, dict[str, float]] = {
        vendor: defaultdict(float) for vendor in vendor_columns.values()
    }
    quoted_totals: dict[str, float] = defaultdict(float)
    unmatched: list[str] = []
    technical_rows: list[VendorQuoteRow] = []
    technical_header = _technical_columns(ws)
    technical_columns = technical_header[1] if technical_header else {}
    for row_index in range(header_row + 1, ws.max_row + 1):
        description = _text(ws.cell(row_index, description_column).value)
        if not description:
            continue
        values = {
            vendor: _number(ws.cell(row_index, column).value)
            for column, vendor in vendor_columns.items()
        }
        if _key(description) == "total":
            quoted_totals.update({vendor: amount or 0.0 for vendor, amount in values.items()})
            continue
        if not any(amount not in (None, 0.0) for amount in values.values()):
            continue
        quantity_value = (
            ws.cell(row_index, technical_columns["quantity"]).value
            if "quantity" in technical_columns
            else None
        )
        unit_value = (
            ws.cell(row_index, technical_columns["unit"]).value
            if "unit" in technical_columns
            else None
        )
        quantity, unit = _quantity_and_unit(quantity_value, unit_value)
        technical_rows.append(
            VendorQuoteRow(
                description=description,
                amounts_usd={vendor: amount for vendor, amount in values.items() if amount is not None},
                construction_scope=(
                    _text(ws.cell(row_index, technical_columns["construction_scope"]).value)
                    if "construction_scope" in technical_columns
                    else ""
                ),
                specification=(
                    _text(ws.cell(row_index, technical_columns["specification"]).value)
                    if "specification" in technical_columns
                    else ""
                ),
                brand=(
                    _text(ws.cell(row_index, technical_columns["brand"]).value)
                    if "brand" in technical_columns
                    else ""
                ),
                quantity=quantity,
                unit=unit,
            )
        )
        category = mappings.get(_key(description))
        if not category:
            unmatched.append(description)
            continue
        for vendor, amount in values.items():
            if amount is not None:
                category_totals[vendor][category] += amount

    for vendor, categories in category_totals.items():
        calculated = sum(categories.values())
        quoted = quoted_totals.get(vendor)
        if quoted and abs(calculated - quoted) > 0.01:
            raise WorkbookValidationError(
                f"{vendor} mapped category total {calculated:.2f} does not match quoted total {quoted:.2f}."
            )
        quoted_totals[vendor] = quoted or calculated
    return category_totals, quoted_totals, sorted(set(unmatched)), technical_rows


def _evaluate_structured(
    data: StructuredBudgetInput,
    source_filename: str,
    report_id: str,
) -> BudgetEvaluation:
    area = data.construction_area_sqm
    if area <= 0:
        raise WorkbookValidationError("Construction area in Sqm is missing or invalid.")
    if not data.per_budget or not data.qs_estimate or not data.historical_unit_rates or not data.mappings:
        raise WorkbookValidationError("One or more required budget sources contain no usable data.")
    if not data.vendor_rows:
        raise WorkbookValidationError("Vendor quotation contains no usable priced rows.")

    mappings = {_key(source): target for source, target in data.mappings.items()}
    category_totals: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    unmatched: list[str] = []
    for row in data.vendor_rows:
        category = mappings.get(_key(row.description))
        if not category:
            unmatched.append(row.description)
            continue
        for vendor, amount in row.amounts_usd.items():
            category_totals[vendor][category] += amount

    quoted_totals = dict(data.vendor_quoted_totals)
    for vendor, categories in category_totals.items():
        calculated = sum(categories.values())
        quoted = quoted_totals.get(vendor)
        if quoted and abs(calculated - quoted) > 0.01:
            raise WorkbookValidationError(
                f"{vendor} mapped category total {calculated:.2f} does not match quoted total {quoted:.2f}."
            )
        quoted_totals[vendor] = quoted or calculated

    return _build_evaluation(
        project_name=data.project_name,
        source_filename=source_filename,
        report_id=report_id,
        area=area,
        per_budget=data.per_budget,
        qs_estimate=data.qs_estimate,
        historical=data.historical_unit_rates,
        vendor_categories=category_totals,
        vendor_totals=quoted_totals,
        unmatched=sorted(set(unmatched)),
        scope_comparisons=_scope_comparisons(data),
    )


def evaluate_structured_input(
    data: StructuredBudgetInput,
    source_filename: str,
    content_digest: bytes,
) -> BudgetEvaluation:
    return _evaluate_structured(
        data,
        source_filename,
        hashlib.sha256(content_digest).hexdigest()[:16],
    )


def _build_evaluation(
    *,
    project_name: str,
    source_filename: str,
    report_id: str,
    area: float,
    per_budget: dict[str, float],
    qs_estimate: dict[str, float],
    historical: dict[str, float],
    vendor_categories: dict[str, dict[str, float]],
    vendor_totals: dict[str, float],
    unmatched: list[str],
    scope_comparisons: list[ScopeComparison] | None = None,
) -> BudgetEvaluation:
    per_lookup = {_key(category): amount for category, amount in per_budget.items()}
    qs_lookup = {_key(category): amount for category, amount in qs_estimate.items()}
    historical_lookup = {_key(category): amount for category, amount in historical.items()}
    display_names = {
        _key(category): category
        for source in (per_budget, qs_estimate, historical)
        for category in source
    }
    normalized_categories = {
        vendor: {
            display_names.get(_key(category), category): amount
            for category, amount in categories.items()
        }
        for vendor, categories in vendor_categories.items()
    }
    per_total = sum(per_lookup.values())
    qs_total = sum(qs_lookup.values())
    historical_weighted = sum(historical_lookup.values())
    warnings: list[str] = []
    if unmatched:
        warnings.append(
            f"{len(unmatched)} priced Vendor Quotation item(s) were not mapped by MAPPING RULES."
        )

    vendors: list[VendorAssessment] = []
    for vendor, categories in normalized_categories.items():
        total = vendor_totals[vendor]
        total_variance = total / per_total - 1 if per_total else 0.0
        qs_comparisons: list[Comparison] = []
        historical_comparisons: list[Comparison] = []
        for category, amount in sorted(categories.items()):
            category_key = _key(category)
            unit_rate = amount / area
            qs_amount = qs_lookup.get(category_key)
            if qs_amount is None or qs_amount == 0:
                qs_comparisons.append(
                    Comparison(
                        category=category,
                        vendor_amount_usd=amount,
                        vendor_unit_rate_usd_sqm=unit_rate,
                        decision="benchmark_missing",
                        status_color="gray",
                    )
                )
            else:
                variance = amount / qs_amount - 1
                decision = _decision(variance, 0.05, 0.15)
                qs_comparisons.append(
                    Comparison(
                        category=category,
                        vendor_amount_usd=amount,
                        vendor_unit_rate_usd_sqm=unit_rate,
                        benchmark_amount_usd=qs_amount,
                        benchmark_unit_rate_usd_sqm=qs_amount / area,
                        variance_pct=variance,
                        decision=decision,
                        status_color=_status_color(decision),
                    )
                )

            historical_rate = historical_lookup.get(category_key)
            if historical_rate is None or historical_rate == 0:
                historical_comparisons.append(
                    Comparison(
                        category=category,
                        vendor_amount_usd=amount,
                        vendor_unit_rate_usd_sqm=unit_rate,
                        decision="benchmark_missing",
                        status_color="gray",
                    )
                )
            else:
                variance = unit_rate / historical_rate - 1
                decision = _decision(variance, 0.05, 0.15)
                historical_comparisons.append(
                    Comparison(
                        category=category,
                        vendor_amount_usd=amount,
                        vendor_unit_rate_usd_sqm=unit_rate,
                        benchmark_unit_rate_usd_sqm=historical_rate,
                        variance_pct=variance,
                        decision=decision,
                        status_color=_status_color(decision),
                    )
                )

        total_decision = _total_decision(total_variance)
        vendors.append(
            VendorAssessment(
                vendor=vendor,
                total_amount_usd=total,
                unit_rate_usd_sqm=total / area,
                per_budget_variance_pct=total_variance,
                per_budget_decision=total_decision,
                per_budget_status_color=_status_color(total_decision),
                recommended_value_usd=min(total, per_total),
                potential_savings_usd=max(total - per_total, 0.0),
                qs_comparisons=qs_comparisons,
                historical_comparisons=historical_comparisons,
            )
        )

    return BudgetEvaluation(
        report_id=report_id,
        project_name=project_name,
        source_filename=source_filename,
        construction_area_sqm=area,
        per_budget_total_usd=per_total,
        per_budget_unit_rate_usd_sqm=per_total / area,
        qs_total_usd=qs_total,
        qs_unit_rate_usd_sqm=qs_total / area,
        historical_weighted_unit_rate_usd_sqm=historical_weighted,
        vendors=vendors,
        unmatched_vendor_items=unmatched,
        scope_comparisons=scope_comparisons or [],
        warnings=warnings,
    )


def evaluate_workbook(content: bytes, filename: str) -> BudgetEvaluation:
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise WorkbookValidationError(f"Unable to open Excel workbook: {exc}") from exc

    sheet_names = {_key(name): name for name in workbook.sheetnames}
    missing = sorted(name for name in REQUIRED_SHEETS if _key(name) not in sheet_names)
    if missing:
        raise WorkbookValidationError(f"Missing required tabs: {', '.join(missing)}")

    def sheet(name: str) -> Worksheet:
        return workbook[sheet_names[_key(name)]]

    summary = sheet("Summary")
    area = _find_label_number(summary, "Construction Area (Sqm)")
    if not area or area <= 0:
        sqft = _find_label_number(summary, "Construction Area (Sqft)")
        area = sqft * 0.09290304 if sqft else None
    if not area or area <= 0:
        raise WorkbookValidationError("Construction area in Sqm is missing or invalid.")

    per_budget = _two_column_table(sheet("PER Budget"), total_prefixes=("total",))
    qs_estimate = _two_column_table(sheet("QS Estimation"), total_prefixes=("total",))
    historical = _two_column_table(
        sheet("Historical Unit Rates"), total_prefixes=("weighted average",)
    )
    mappings = _mapping_table(sheet("MAPPING RULES"))
    if not per_budget or not qs_estimate or not historical or not mappings:
        raise WorkbookValidationError("One or more budget benchmark tabs contain no usable data.")

    vendor_categories, vendor_totals, unmatched, vendor_rows = _vendor_table(
        sheet("Vendor Quotation"), mappings
    )
    cover_name = sheet_names.get("cover")
    project_name = (
        _text(workbook[cover_name]["A1"].value)
        if cover_name
        else _text(summary["A1"].value)
    ) or Path(filename).stem
    scope_data = StructuredBudgetInput(
        project_name=project_name,
        construction_area_sqm=area,
        per_budget=per_budget,
        qs_estimate=qs_estimate,
        historical_unit_rates=historical,
        mappings=mappings,
        vendor_rows=vendor_rows,
        vendor_quoted_totals=dict(vendor_totals),
        qs_rows=_qs_requirement_rows(sheet("QS Estimation")),
    )
    return _build_evaluation(
        report_id=hashlib.sha256(content).hexdigest()[:16],
        project_name=project_name,
        source_filename=Path(filename).name,
        area=area,
        per_budget=per_budget,
        qs_estimate=qs_estimate,
        historical=historical,
        vendor_categories=vendor_categories,
        vendor_totals=vendor_totals,
        unmatched=unmatched,
        scope_comparisons=_scope_comparisons(scope_data),
    )
